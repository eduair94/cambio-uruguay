"""Extract only five exact, labelled concepts from the original public MSP XLSX.

Requires Python 3.10+ and openpyxl. No network or spreadsheet modification.
Only replaces the GENERATED block in app/utils/mutualistaCosts.ts. Fail closed
on a different source file, missing concepts, unknown regime headers, wrong
dates or nonnumeric prices. See docs/app/MUTUALISTA_COSTS.md before refreshing.
"""
from pathlib import Path
from datetime import datetime
import argparse
import hashlib
import json
import math
import re
import unicodedata
import openpyxl

EXPECTED_SOURCE_SHA256 = '1c189260422ff5a131bbb2e87d862b1c0226974378d2d402eea5337a868df821'

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source', type=Path, required=True, help='Local original MSP XLSX to read.')
parser.add_argument('--report', type=Path, help='Optional output path for a JSON extraction report.')
args = parser.parse_args()
source_file = args.source.expanduser().resolve()
target = Path(__file__).resolve().parents[1] / 'utils/mutualistaCosts.ts'
report_path = args.report.expanduser().resolve() if args.report else None
if not source_file.is_file():
    parser.error('--source must name an existing local XLSX file')
if report_path in (source_file, target):
    parser.error('--report must not overwrite the source XLSX or generated TypeScript module')
if report_path and not report_path.parent.is_dir():
    parser.error('--report parent directory must already exist')
sha = hashlib.sha256(source_file.read_bytes()).hexdigest()
if sha != EXPECTED_SOURCE_SHA256:
    parser.error(
        'Source SHA-256 does not match the reviewed July 2026 edition. '
        'Review the new publication and update the edition guards before importing it.'
    )
wb = openpyxl.load_workbook(source_file, read_only=True, data_only=True)
concepts = [
    'TICKET DE MEDICAMENTOS - GENERAL',
    'CONSULTA NO URGENTE CONSULTORIO: MEDICINA GENERAL',
    'CONSULTA NO URGENTE CONSULTORIO: OTRAS ESPECIALIDADES',
    'CONSULTA URGENCIA CENTRALIZADA',
    'CONSULTA NO URGENCIA DOMICILIO',
]
groups = {'NO FONASA': 'no-fonasa', 'FONASA': 'fonasa', 'SANIDAD POLICIAL': 'sanidad-policial'}

def normalized(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().upper()

def numeric(value, cell):
    if value is None or value == '':
        return None
    if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value) or value < 0:
        raise ValueError(f'Invalid numeric value in {cell}: {value!r}')
    return value

def slug(value):
    text = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode().lower()
    return re.sub('[^a-z0-9]+', '-', text).strip('-')

raw = []
report = {'sourceSha256': sha, 'concepts': concepts, 'institutions': [], 'anomalies': []}
for ws in wb.worksheets:
    # The source formats thousands of blank rows; only rows with the selected
    # literal labels are relevant. Keep the data-cached values, not formulas.
    rows = list(ws.iter_rows(max_row=300, values_only=True))
    if not isinstance(rows[6][1], datetime) or rows[6][1].strftime('%Y-%m-%d') != '2026-07-01':
        raise ValueError(f'Wrong source date in {ws.title}!B7')
    if normalized(rows[10][5]) != 'VALOR MÁXIMO AUTORIZADO':
        raise ValueError(f'Changed maximum header in {ws.title}!F11')
    selected = []
    for concept in concepts:
        matches = [(index + 1, row) for index, row in enumerate(rows) if normalized(row[0]) == concept]
        if len(matches) != 1:
            raise ValueError(f'Expected exactly one {concept!r} in {ws.title}, found {len(matches)}')
        if any(value not in (None, '') for value in matches[0][1][1:5]):
            raise ValueError(f'Selected concept has an extra qualifier in {ws.title}: {concept}')
        selected.append(matches[0])
    maxima = [numeric(row[5], f'{ws.title}!F{num}') for num, row in selected]
    if any(value is not None and value > 880 for value in maxima):
        raise ValueError(f'Selected concept exceeds source general ceiling: {ws.title}')
    source_rows = [[num, str(row[0]).strip()] for num, row in selected]
    columns = []
    group = None
    for col_index in range(6, ws.max_column):
        header_group = rows[8][col_index]
        if header_group not in (None, ''):
            group = groups.get(normalized(header_group))
            if group is None:
                raise ValueError(f'Unknown affiliation {header_group!r} in {ws.title}')
        label = rows[10][col_index]
        if label in (None, ''):
            continue
        if group is None or 'PRECIO' not in normalized(label):
            raise ValueError(f'Unexpected price column in {ws.title}: {label!r}')
        condition = rows[9][col_index]
        # Auxiliary numeric percentages do not identify a named convention.
        # The main header (e.g. CASMU's explicit discount) is always preserved.
        condition = condition.strip() if isinstance(condition, str) and condition.strip() else None
        column = openpyxl.utils.get_column_letter(col_index + 1)
        values = [numeric(row[col_index], f'{ws.title}!{column}{num}') for num, row in selected]
        for concept_index, (value, maximum) in enumerate(zip(values, maxima)):
            if value is not None and maximum is not None and value > maximum + 0.005:
                report['anomalies'].append({'sheet': ws.title, 'cell': f'{column}{source_rows[concept_index][0]}',
                                            'affiliation': group, 'concept': concepts[concept_index],
                                            'amount': value, 'publishedMaximum': maximum})
        columns.append([column, group, str(label).strip(), condition, values])
    if not {'fonasa', 'no-fonasa'}.issubset(set(column[1] for column in columns)):
        raise ValueError(f'Missing affiliation group in {ws.title}')
    raw.append([slug(ws.title), str(rows[5][1]).strip(), ws.title, '2026-07-01', maxima, source_rows, columns])
    values = [value for column in columns for value in column[4]]
    report['institutions'].append({'id': slug(ws.title), 'name': rows[5][1], 'sheet': ws.title,
                                   'columns': len(columns), 'groups': sorted(set(col[1] for col in columns)),
                                   'nulls': sum(value is None for value in values),
                                   'zeros': sum(value == 0 for value in values),
                                   'maxima': maxima, 'sourceRows': source_rows})
wb.close()
if len(raw) != 34 or len({item[0] for item in raw}) != 34:
    raise ValueError('Unexpected institution coverage')

def js(value):
    return json.dumps(value, ensure_ascii=False, separators=(', ', ': '))

page = 'https://www.gub.uy/ministerio-salud-publica/datos-y-estadisticas/datos/precios-tickets-ordenes-instituciones-asistencia-medica-colectiva-iamc-julio-2026'
download = 'https://www.gub.uy/ministerio-salud-publica/sites/ministerio-salud-publica/files/2026-09/Precios%20Tasas%20Moderadoras%20-%20Julio%202026.xlsx'
metadata = {'pageUrl': page, 'downloadUrl': download, 'publishedAt': '2026-09-03',
            'effectiveFrom': '2026-07-01', 'verifiedAt': '2026-09-14', 'sha256': sha,
            'ivaRate': 0.1, 'medicineStampUyu': 44, 'diagnosticStampUyu': 170}
lines = ['export const MUTUALISTA_COST_SOURCE = ' + js(metadata) + ' as const', '',
         '// Compact source tuples: preserve column order and unrounded numeric values.',
         '// prettier-ignore', 'const RAW_INSTITUTIONS: readonly RawInstitution[] = [']
for institution in raw:
    lines.append('  [')
    lines.append('    ' + ', '.join(js(value) for value in institution[:4]) + ',')
    lines.append('    ' + js(institution[4]) + ',')
    lines.append('    ' + js(institution[5]) + ',')
    lines.append('    [')
    lines.extend('      ' + js(column) + ',' for column in institution[6])
    lines.extend(['    ],', '  ],'])
lines.append(']')
before = target.read_text(encoding='utf-8')
start, end = '// BEGIN GENERATED MSP DATA', '// END GENERATED MSP DATA'
if before.count(start) != 1 or before.count(end) != 1:
    raise ValueError('Cannot find unique generation markers')
after = before.split(start)[0] + start + '\n' + '\n'.join(lines) + '\n' + end + before.split(end)[1]
target.write_text(after, encoding='utf-8', newline='\n')
if report_path:
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'institutions': len(raw), 'columns': sum(len(item[6]) for item in raw), 'sourceSha256': sha,
                  'target': str(target)}, ensure_ascii=True))
